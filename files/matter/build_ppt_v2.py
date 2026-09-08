# -*- coding: utf-8 -*-
"""基于 PAA.pptx 模板生成《Matter 设备认证验签详解 v2》内部分享 PPT。
内容页全部从模板 slide2 的版式 (slideLayout2) 派生，保留公司 logo 等装饰。
v2 修正：标题遮挡、空单元格、页脚碰撞、封面副标题换行、OID 树连线、留白平衡等视觉 QA 问题。
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

SRC = r"D:\hrf\h\files\matter\PAA.pptx"
DST = r"D:\hrf\h\files\matter\PAA_PAI_DAC验签详解_v2_内部分享.pptx"

# ---------- palette ----------
RED    = "C00000"   # 主色（标题/PAA）
DARK   = "333333"
ORANGE = "D9482F"   # PAI（加深以提升白字对比度）
BLUE   = "1F6FC5"   # CD/强调
GRAY   = "808080"
TBL_HDR = RED
TBL_ALT = "F2F2F2"
CODE_BG = "F5F5F5"
RED_BG   = "FBEAEA"
BLUE_BG  = "EAF2FB"
GREEN_BG = "EAF4EA"
GREEN    = "2E7D32"
WHITE   = "FFFFFF"

YH = "微软雅黑"

BODY_X, BODY_Y, BODY_W, BODY_H = 0.51, 0.98, 12.31, 6.05


def rgb(h):
    return RGBColor.from_string(h)


def style_run(run, size=14, bold=False, color=DARK, font=YH, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = rgb(color)
    f.name = font
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", "zh-CN")
    rPr.set("altLang", "en-US")
    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = rPr.makeelement(qn("a:ea"), {})
        rPr.append(ea)
    ea.set("typeface", font)


def para(tf, text, size=14, bold=False, color=DARK, align=PP_ALIGN.LEFT,
         before=2, after=2, bullet=None, line=None, font=YH, first=False, italic=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    if bullet:
        pPr = p._p.get_or_add_pPr()
        buFont = pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"})
        buChar = pPr.makeelement(qn("a:buChar"), {"char": bullet})
        pPr.append(buFont)
        pPr.append(buChar)
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=bold, color=color, font=font, italic=italic)
    return p


def add_text(slide, x, y, w, h, wrap=True, anchor=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.01)
    tf.margin_bottom = Inches(0.01)
    if anchor:
        tf.vertical_anchor = anchor
    return tf


def add_box(slide, x, y, w, h, fill=None, line=None, round_=True, radius=0.10, line_w=0.75):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h))
    if round_:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = rgb(fill)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = rgb(line)
        shp.line.width = Pt(line_w)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.08)
    tf.margin_right = Inches(0.08)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return shp


def add_code(slide, x, y, w, h, lines, size=10.5, title=None):
    box = add_box(slide, x, y, w, h, fill=CODE_BG, line="DDDDDD", radius=0.05)
    tf = box.text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = Inches(0.12)
    tf.margin_top = Inches(0.08)
    first = True
    if title:
        para(tf, title, size=10, bold=True, color=GRAY, first=True, after=4)
        first = False
    for ln in lines:
        para(tf, ln, size=size, color="1A1A1A", font="Consolas", first=first,
             before=0, after=0, line=1.0)
        first = False
    return box


def add_arrow(slide, x1, y1, x2, y2, color=GRAY, w=1.5):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(w)
    conn.shadow.inherit = False
    ln = conn.line._get_or_add_ln()
    tail = ln.makeelement(qn("a:tailEnd"), {"type": "triangle", "w": "med", "len": "med"})
    ln.append(tail)
    return conn


def add_line(slide, x1, y1, x2, y2, color=GRAY, w=1.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                      Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    conn.line.color.rgb = rgb(color)
    conn.line.width = Pt(w)
    conn.shadow.inherit = False
    if dash:
        ln = conn.line._get_or_add_ln()
        d = ln.makeelement(qn("a:prstDash"), {"val": dash})
        ln.append(d)
    return conn


def set_table_style(table):
    table.first_row = False
    table.horz_banding = False


def fill_cell(cell, text, size=11, bold=False, color=DARK, fill=None,
              align=PP_ALIGN.LEFT, font=YH):
    if fill:
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Inches(0.06)
    cell.margin_right = Inches(0.06)
    cell.margin_top = Inches(0.02)
    cell.margin_bottom = Inches(0.02)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    for r in list(p.runs):
        r._r.getparent().remove(r._r)
    run = p.add_run()
    run.text = text
    style_run(run, size=size, bold=bold, color=color, font=font)


def add_table(slide, x, y, w, headers, rows, col_ratios=None, header_size=11,
              body_size=10.5, row_h=0.32, header_h=0.34, aligns=None):
    n_rows = len(rows) + 1
    n_cols = len(headers)
    gfx = slide.shapes.add_table(n_rows, n_cols, Inches(x), Inches(y), Inches(w),
                                 Inches(header_h + row_h * len(rows)))
    table = gfx.table
    set_table_style(table)
    if col_ratios:
        total = sum(col_ratios)
        for i, ratio in enumerate(col_ratios):
            table.columns[i].width = Emu(int(Inches(w) * ratio / total))
    table.rows[0].height = Inches(header_h)
    for j, htext in enumerate(headers):
        fill_cell(table.cell(0, j), htext, size=header_size, bold=True, color=WHITE,
                  fill=TBL_HDR, align=PP_ALIGN.CENTER)
    for i, row in enumerate(rows, start=1):
        table.rows[i].height = Inches(row_h)
        bg = TBL_ALT if i % 2 == 0 else None
        for j, ctext in enumerate(row):
            al = (aligns[j] if (aligns and j < len(aligns)) else PP_ALIGN.LEFT)
            fill_cell(table.cell(i, j), ctext, size=body_size, fill=bg, align=al)
    return table


def content_slide(prs, layout, title):
    slide = prs.slides.add_slide(layout)
    ph = slide.shapes.title
    ph.text = title
    return slide


def section_label(slide, text, x=None, y=None, color=RED, w=None):
    tf = add_text(slide, BODY_X if x is None else x,
                  BODY_Y - 0.05 if y is None else y,
                  BODY_W if w is None else w, 0.34)
    para(tf, text, size=15, bold=True, color=color, first=True, after=4)


# ============================================================
# 各页构建函数
# ============================================================

def s_toc(prs, layout):
    slide = content_slide(prs, layout, "目录")
    items_l = [
        ("01", "本版修订说明", "以 Matter 1.5 规范 + SDK 源码逐条核对"),
        ("02", "验签在配网流程中的位置", "从 PASE 会话到 AddNOC"),
        ("03", "三张证书与一份声明", "PAA / PAI / DAC / CD 各自的角色"),
        ("04", "VID/PID 编码细节", "Matter OID、preferred/fallback"),
        ("05", "验签全流程", "SDK 真实执行的 12 步 + 真机日志解读"),
    ]
    items_r = [
        ("06", "CD 深度解析", "CMS 结构、6 把钥匙、交叉校验矩阵"),
        ("07", "常见疑问 FAQ", "PAI 能否带 PID 列表？PID 不在 CD？"),
        ("08", "设备端落地（Silabs）", "Provider 接口、NVM3、私钥不出安全区"),
        ("09", "错误码与排查", "AttestationVerificationResult 速查"),
        ("10", "生产 Checklist", "从开发到量产的检查清单"),
    ]
    for col, items in ((0, items_l), (1, items_r)):
        x = BODY_X + col * 6.28
        for i, (num, t, sub) in enumerate(items):
            y = 1.15 + i * 1.12
            add_box(slide, x, y, 0.62, 0.62, fill=RED, radius=0.28)
            tfn = add_text(slide, x, y + 0.10, 0.62, 0.42)
            para(tfn, num, size=16, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
            tft = add_text(slide, x + 0.78, y - 0.02, 5.2, 0.42)
            para(tft, t, size=16, bold=True, color=DARK, first=True)
            tfs = add_text(slide, x + 0.78, y + 0.38, 5.2, 0.36)
            para(tfs, sub, size=11, color=GRAY, first=True)


def s_revisions(prs, layout):
    slide = content_slide(prs, layout, "本版修订了什么")
    headers = ["#", "旧版说法", "正确说法（本版）"]
    rows = [
        ["1", "PAA 的 pathLen 可为 0 或 1，pathLen:0 = PAA 直接签发 DAC",
         "pathLen 缺省或必须为 1；v1.5 链路必须恰好 3 层（PAA→PAI→DAC）"],
        ["2", "PAI KeyUsage = keyCertSign",
         "必须 keyCertSign + cRLSign 两位同时置 1（digitalSignature 可选）"],
        ["3", "CD 的 security_level 是 32-bit",
         "8-bit（SDK：uint8_t SecurityLevel）"],
        ["4", "CD 校验基准是 DAC 的 VID/PID",
         "第一基准是 Basic Information 集群，之后才交叉 DAC/PAI"],
        ["5", "AttestationResponse 的 Tag2=DAC、Tag3=PAI",
         "DAC/PAI 走 CertificateChainRequest 单独获取（1=DAC、2=PAI）"],
        ["6", "CD 校验在 attestation 签名验证之前",
         "SDK 真实顺序：签名→找 PAA→链验证→nonce→CD"],
    ]
    add_table(slide, BODY_X, 1.05, BODY_W, headers, rows,
              col_ratios=[0.5, 4.3, 7.0], body_size=11, row_h=0.72, header_h=0.36)
    tf = add_text(slide, BODY_X, 6.25, BODY_W, 0.4)
    para(tf, "依据：Matter Core Spec R1.5（23-27349-009）逐条核对 + Silicon Labs Matter SDK 源码逐行核对",
         size=11, color=GRAY, italic=True, first=True)


def s_new_contents(prs, layout):
    slide = content_slide(prs, layout, "本版新增 / 加强了什么")
    section_label(slide, "在旧版正确内容（证书链架构、SKID/AKID、签名算法、配网时序、日志分析）之上，新增：")
    items = [
        ("SDK 真实 12 步验证序列", "对齐 DefaultDACVerifier 源码实际执行顺序，每步带错误码"),
        ("CD 验签公钥全表", "SDK 内置 6 把钥匙：1 把测试 + 官方 001~005，含完整 Key ID"),
        ("dac_origin（ODM 贴牌）校验矩阵", "Basic Info 报品牌身份、DAC 报原厂身份、CD 把两者绑定"),
        ("PAI.PID 与 CD 列表的联动规则", "PAI 带 PID 时该 PID 也必须在 CD.product_id_array 内"),
        ("设备端落地", "Provider 接口、私钥不出安全区、NOCSR 复用 DAC 私钥"),
        ("DCL 吊销机制", "为什么 Matter 不用 OCSP、cRLDistributionPoints 为何被忽略"),
    ]
    for i, (t, d) in enumerate(items):
        y = 1.62 + i * 0.86
        add_box(slide, BODY_X, y, 0.14, 0.6, fill=RED, radius=0.5)
        tft = add_text(slide, BODY_X + 0.30, y - 0.03, 8.6, 0.4)
        para(tft, t, size=14.5, bold=True, color=DARK, first=True)
        tfd = add_text(slide, BODY_X + 0.30, y + 0.30, 8.6, 0.36)
        para(tfd, d, size=11.5, color=GRAY, first=True)
    add_box(slide, 9.7, 1.62, 3.1, 4.15, fill=BLUE_BG)
    tf = add_text(slide, 9.92, 1.78, 2.66, 3.9)
    para(tf, "核对基准", size=14, bold=True, color=BLUE, first=True, after=8)
    para(tf, "① Matter Core Spec R1.5（23-27349-009）逐条核对", size=11.5, after=8)
    para(tf, "② Silicon Labs Matter SDK 源码逐行核对", size=11.5, after=8)
    para(tf, "③ chip-tool 实测日志对照", size=11.5, after=8)
    para(tf, "任何与两处权威源不符的表述，均已在本版修正。", size=11.5, bold=True, color=BLUE)


def s_flow_position(prs, layout):
    slide = content_slide(prs, layout, "验签发生在配网流程的哪里")
    left_x = BODY_X + 0.25
    steps = [
        ("1", "PASE 会话建立", "BLE → SPAKE2+（PBKDFParamRequest → Pake1~3）；派生 16 字节 AttestationChallenge"),
        ("2", "读取设备信息", "ReadCommissioningInfo：Basic Info 集群自报 vendorId=0x149A productId=0x3005"),
        ("3", "ArmFailSafe + ConfigRegulatory", "布 60 秒失败保险；地区配置（条款确认可跳过）"),
        ("4", "请求 PAI / DAC 证书", "CertificateChainRequest type=2/1 → CertificateChainResponse（470B / 481B DER）"),
        ("5", "设备认证", "AttestationRequest { nonce 32B } → elements 423B + AttestationSignature 64B"),
        ("6", "本地验证 12 步 + 吊销检查", "kAttestationVerification（后两页展开）→ kAttestationRevocationCheck"),
        ("7", "CSR → NOC → 网络 → CASE → Complete", "NOCSR 复用 DAC 私钥；装 NOC 后 Thread 配网，DNS-SD 发现 → CASE → CommissioningComplete"),
    ]
    y0 = 1.08
    for i, (num, t, d) in enumerate(steps):
        y = y0 + i * 0.82
        if i < len(steps) - 1:
            add_arrow(slide, left_x + 0.21, y + 0.62, left_x + 0.21, y + 0.84, color="888888", w=2)
        add_box(slide, left_x, y, 0.42, 0.42, fill=RED, radius=0.5)
        tfn = add_text(slide, left_x, y + 0.06, 0.42, 0.3)
        para(tfn, num, size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
        add_box(slide, left_x + 0.58, y - 0.06, 11.35, 0.66, fill="F7F7F7", line="E4E4E4", radius=0.14)
        tft = add_text(slide, left_x + 0.74, y - 0.03, 11.0, 0.3)
        para(tft, t, size=12.5, bold=True, color=RED, first=True)
        tfd = add_text(slide, left_x + 0.74, y + 0.24, 11.0, 0.34)
        para(tfd, d, size=10.5, color=DARK, first=True)
    tf = add_text(slide, BODY_X, 6.95, 8.6, 0.34)
    para(tf, "阶段名来自 SDK CommissioningDelegate.h；数值取自真机日志 commissioning-raspi-log.md",
         size=10.5, color=GRAY, italic=True, first=True)


def s_separate(prs, layout):
    slide = content_slide(prs, layout, "证书链与 CD 为什么分开传")
    add_box(slide, BODY_X, 1.05, 6.0, 2.55, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, BODY_X + 0.2, 1.18, 5.6, 0.4)
    para(tf, "证书链：静态身份", size=16, bold=True, color=ORANGE, first=True)
    tf2 = add_text(slide, BODY_X + 0.2, 1.64, 5.6, 1.9)
    for i, t in enumerate([
        "CertificateChainRequest / Response 命令",
        "内容静态不变 → 取一次可缓存复用",
        "每张 ≤ 600 字节 DER（§6.1.3）",
        "DAC / PAI 证书本体，证明\u201c我是谁\u201d",
    ]):
        para(tf2, t, size=12.5, bullet="▪", first=(i == 0), after=5)
    add_box(slide, BODY_X + 6.31, 1.05, 6.0, 2.55, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, BODY_X + 6.51, 1.18, 5.6, 0.4)
    para(tf, "CD + nonce + 时间戳：动态证据", size=16, bold=True, color=BLUE, first=True)
    tf2 = add_text(slide, BODY_X + 6.51, 1.64, 5.6, 1.9)
    for i, t in enumerate([
        "AttestationRequest / Response 命令",
        "绑定本次配网 + 当前安全会话",
        "nonce 防重放、challenge 防挪用",
        "证明\u201c型号已认证且是它本人在回答\u201d",
    ]):
        para(tf2, t, size=12.5, bullet="▪", first=(i == 0), after=5)
    add_box(slide, BODY_X, 3.85, BODY_W, 2.0, fill=RED_BG, line=None)
    tf3 = add_text(slide, BODY_X + 0.25, 4.0, BODY_W - 0.5, 1.75)
    para(tf3, "⚠ 纠正一个常见误解", size=14, bold=True, color=RED, first=True, after=4)
    para(tf3, "AttestationResponse 里并没有 DAC/PAI 证书！它只有 2 个字段：", size=13, after=3)
    para(tf3, "AttestationElements [0]（≤900B，内嵌 CD+nonce+时间戳）    AttestationSignature [1]（64B r||s）",
         size=13, bold=True, color=RED, after=3)
    para(tf3, "证书本体只能从 CertificateChainResponse 拿（规范 §11.18.6.2 / §11.18.4.6）",
         size=12, color=GRAY)


def s_analogy(prs, layout):
    slide = content_slide(prs, layout, "用一个类比建立直觉")
    headers = ["现实世界", "Matter 世界", "一句话"]
    rows = [
        ["公安部 / 根身份证签发机构", "PAA", "信任锚点，Commissioner 只信任名单上的 PAA"],
        ["户籍派出所 / 分局", "PAI", "代表一个厂商（VID），或一条产品线（VID+PID）"],
        ["个人身份证", "DAC", "每台设备唯一，\u201c我是谁\u201d"],
        ["学历 / 学位认证报告", "CD", "CSA 联盟签发：\u201c这个 (VID,PID) 通过了认证\u201d"],
    ]
    add_table(slide, BODY_X, 1.15, BODY_W, headers, rows,
              col_ratios=[3.2, 2.2, 6.5], body_size=13, row_h=0.78, header_h=0.4,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.LEFT])
    add_box(slide, BODY_X, 5.0, BODY_W, 1.25, fill=BLUE_BG)
    tf = add_text(slide, BODY_X + 0.3, 5.14, BODY_W - 0.6, 1.0)
    para(tf, "一句话记住", size=13, bold=True, color=BLUE, first=True, after=4)
    para(tf, "证书链证明\u201c我是谁、谁为我背书\u201d；CD 证明\u201c我这个型号通过了官方认证\u201d。两者缺一不可，配网时同时校验。",
         size=15, bold=True, color=DARK)


def s_pki(prs, layout):
    slide = content_slide(prs, layout, "PKI 层级与三条硬性规定")
    bx = BODY_X + 0.1
    add_box(slide, bx, 1.15, 3.5, 1.05, fill=RED)
    tf = add_text(slide, bx, 1.24, 3.5, 0.9)
    para(tf, "PAA  信任根", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=2)
    para(tf, "CA:TRUE，pathLen 缺省或 1 ｜ 自签名", size=10, color="FFDDDD", align=PP_ALIGN.CENTER)
    add_arrow(slide, bx + 1.75, 2.20, bx + 1.75, 2.52, color=DARK, w=2)
    add_box(slide, bx, 2.52, 3.5, 1.05, fill=ORANGE)
    tf = add_text(slide, bx, 2.61, 3.5, 0.9)
    para(tf, "PAI  中间证书", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=2)
    para(tf, "CA:TRUE，pathLen=0 ｜ 必含 VID", size=10, color="FFE8E0", align=PP_ALIGN.CENTER)
    add_arrow(slide, bx + 1.75, 3.57, bx + 1.75, 3.89, color=DARK, w=2)
    add_box(slide, bx, 3.89, 3.5, 1.05, fill=DARK)
    tf = add_text(slide, bx, 3.98, 3.5, 0.9)
    para(tf, "DAC  设备证书", size=15, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True, after=2)
    para(tf, "CA:FALSE ｜ 必含 VID+PID ｜ 私钥不出设备", size=10, color="DDDDDD", align=PP_ALIGN.CENTER)
    tf = add_text(slide, bx, 5.15, 3.9, 0.8)
    para(tf, "真机样例：HOPERF Matter PAA 01 → PAI 01 → DAC（VID 0x1470，DAC 加 PID 0x8006）", size=10, color=GRAY, first=True)
    para(tf, "存放：PAA → Commissioner 信任库 + DCL；PAI/DAC → 设备固件 / 工厂数据", size=10.5, color=GRAY)

    rx = 4.9
    rules = [
        ("硬规定 1", "DAC 必须由 PAI 签发，认证路径长度固定为 2，整链恰好 3 张证书（§6.2.2 / §6.2.3.1）"),
        ("硬规定 2", "PAI 必须归属一个 VID；可再限定一个 PID；服务多个产品就必须不带 PID（§6.2.2.1）"),
        ("硬规定 3", "所有证书 DER ≤ 600 字节；SKID / AKID 固定 20 字节（§6.1.2 / §6.1.3）"),
    ]
    for i, (t, d) in enumerate(rules):
        y = 1.15 + i * 1.1
        add_box(slide, rx, y, 7.9, 0.92, fill="FAFAFA", line="DDDDDD")
        tf = add_text(slide, rx + 0.2, y + 0.07, 7.5, 0.32)
        para(tf, t, size=13, bold=True, color=RED, first=True)
        tf = add_text(slide, rx + 0.2, y + 0.38, 7.5, 0.5)
        para(tf, d, size=11.5, color=DARK, first=True)
    add_box(slide, rx, 4.55, 7.9, 0.85, fill=BLUE_BG)
    tf = add_text(slide, rx + 0.2, 4.66, 7.5, 0.65)
    para(tf, "证书里的 cRLDistributionPoints 可以存在，但 Commissioner 必须忽略 —— 吊销统一走 DCL（§6.2.4）",
         size=11.5, color=DARK, first=True)


def s_vidpid_map(prs, layout):
    slide = content_slide(prs, layout, "谁必须有 VID / PID（全景表）")
    headers = ["位置", "VID", "PID", "说明"]
    rows = [
        ["PAA  Subject", "0 或 1 个", "禁止", "共享型 PAA 常不带 VID；SDK 强制有 PID → kPaaFormatInvalid"],
        ["PAI  Subject", "必须 1 个", "0 或 1 个", "带 PID = 只能服务这一个产品；不带 = 服务该 VID 下所有产品"],
        ["DAC  Subject", "必须 1 个", "必须 1 个", "设备的\u201c身份证号\u201d"],
        ["CD  TLV", "vendor_id ×1", "product_id_array ×1..100", "\u201cPID 列表\u201d只存在于 CD，不在 PAI"],
        ["Basic Information 集群 (0x0028)", "VendorID 属性", "ProductID 属性", "设备运行时自报家门，是 CD 校验的第一基准"],
    ]
    add_table(slide, BODY_X, 1.1, BODY_W, headers, rows,
              col_ratios=[2.9, 1.7, 2.4, 6.6], body_size=11.5, row_h=0.66, header_h=0.38)
    add_box(slide, BODY_X, 5.15, BODY_W, 1.35, fill=RED_BG)
    tf = add_text(slide, BODY_X + 0.25, 5.3, BODY_W - 0.5, 1.1)
    para(tf, "记住两个\u201c只能 0 或 1 个\u201d", size=13, bold=True, color=RED, first=True, after=3)
    para(tf, "PAI / DAC 的 PID 是 X.509 DN 里的单值 RDN，天生不支持列表。想要\u201c一个中间证书覆盖多个型号\u201d，"
             "唯一做法：PAI 不带 PID + CD 的 product_id_array 列出全部已认证 PID。",
         size=12.5, color=DARK)


def s_cert_table(prs, layout):
    slide = content_slide(prs, layout, "证书要求逐项对照（§6.2.2.3 / §6.2.2.4 / §6.2.2.5）")
    headers = ["检查项", "DAC", "PAI", "PAA"]
    rows = [
        ["版本 / 签名算法 / 曲线", "v3 ｜ ecdsa-with-SHA256 ｜ prime256v1",
         "v3 ｜ ecdsa-with-SHA256 ｜ prime256v1", "v3 ｜ ecdsa-with-SHA256 ｜ prime256v1"],
        ["Basic Constraints", "critical, CA:FALSE", "critical, CA:TRUE pathLen=0", "critical, CA:TRUE pathLen 缺省或 1"],
        ["Key Usage", "critical, 仅 digitalSignature", "critical, keyCertSign+cRLSign，±digitalSignature", "critical, keyCertSign+cRLSign，±digitalSignature"],
        ["SKID / AKID", "都必须", "都必须", "SKID 必须；AKID 可选"],
        ["Subject VID / PID", "1 VID + 1 PID", "1 VID + 0..1 PID", "0..1 VID + 0 PID"],
        ["issuer 要求", "= PAI 的 subject（逐字节）", "= PAA 的 subject（逐字节）", "= subject（自签）"],
        ["其他可选扩展", "ExtKeyUsage、AIA、SAN 等", "ExtKeyUsage 等", "ExtKeyUsage、AKID 等"],
    ]
    add_table(slide, BODY_X, 1.1, BODY_W, headers, rows,
              col_ratios=[2.3, 3.2, 3.6, 3.6], body_size=10.5, row_h=0.62, header_h=0.36)
    tf = add_text(slide, BODY_X, 5.95, BODY_W, 0.7)
    para(tf, "注意：三张证书的 Basic Constraints 与 Key Usage 都必须标 critical —— SDK 的格式校验会直接拒绝非 critical 证书。",
         size=12, color=RED, bold=True, first=True)
    para(tf, "v1.4 中\u201cPAA 直接签发 DAC\u201d的两层链写法在 v1.5 已不成立：PAI 是强制存在的。", size=11, color=GRAY)


def s_oid(prs, layout):
    slide = content_slide(prs, layout, "Matter OID 分配（Appendix E, Table 128）")
    add_box(slide, 5.2, 1.05, 2.9, 0.62, fill=DARK)
    tf = add_text(slide, 5.2, 1.15, 2.9, 0.42)
    para(tf, "1.3.6.1.4.1.37244", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    # tree connectors
    add_line(slide, 6.65, 1.67, 6.65, 1.81, color="999999", w=1.25)
    add_line(slide, 3.7, 1.81, 9.7, 1.81, color="999999", w=1.25)
    add_arrow(slide, 3.7, 1.81, 3.7, 1.99, color="999999", w=1.25)
    add_arrow(slide, 9.7, 1.81, 9.7, 1.99, color="999999", w=1.25)
    # .1.x branch
    add_box(slide, BODY_X + 0.4, 2.0, 5.6, 0.5, fill=GRAY)
    tf = add_text(slide, BODY_X + 0.55, 2.06, 5.35, 0.4)
    para(tf, ".1.x = matter-op-cert（运营证书 NOC/ICAC/RCAC 的 DN 属性）", size=11.5, bold=True, color=WHITE, first=True)
    add_box(slide, BODY_X + 0.65, 2.65, 5.35, 1.95, fill="F7F7F7", line="DDDDDD")
    tf = add_text(slide, BODY_X + 0.82, 2.74, 5.05, 1.8)
    lines = [
        ".1.1  matter-node-id             Node ID",
        ".1.2  matter-firmware-signing-id",
        ".1.3  matter-icac-id",
        ".1.4  matter-rcac-id",
        ".1.5  matter-fabric-id",
        ".1.6  matter-noc-cat             CASE 认证标签 CAT",
        ".1.7  matter-vvs-id              VID 验证签名者（1.5 新增）",
    ]
    for i, ln in enumerate(lines):
        para(tf, ln, size=10.5, font="Consolas", first=(i == 0), after=2, line=1.05)
    # .2.x branch
    add_box(slide, BODY_X + 6.4, 2.0, 5.6, 0.5, fill=RED)
    tf = add_text(slide, BODY_X + 6.55, 2.06, 5.35, 0.4)
    para(tf, ".2.x = matter-att-cert（设备认证证书 DN 属性）★ 本档主角", size=11.5, bold=True, color=WHITE, first=True)
    add_box(slide, BODY_X + 6.65, 2.65, 5.35, 1.3, fill=RED_BG)
    tf = add_text(slide, BODY_X + 6.85, 2.78, 5.0, 1.15)
    para(tf, ".2.1  matter-oid-vid      Vendor ID", size=12, font="Consolas", bold=True, color=RED, first=True, after=5)
    para(tf, ".2.2  matter-oid-pid      Product ID", size=12, font="Consolas", bold=True, color=RED)
    add_box(slide, BODY_X, 5.0, BODY_W, 1.0, fill=BLUE_BG)
    tf = add_text(slide, BODY_X + 0.25, 5.12, BODY_W - 0.5, 0.8)
    para(tf, "口诀：.2.x 管认证（VID/PID），.1.x 管运营（NodeID / FabricID / CAT）", size=14.5, bold=True, color=BLUE, first=True, after=2)
    para(tf, "生成证书时 OID 用错 arc，证书过不了 Matter 合规校验。", size=12, color=DARK)
    tf = add_text(slide, BODY_X, 6.25, BODY_W, 0.6)
    para(tf, "DN 属性完整表见规范正文 §6.1.1 Table 83；v1.4 的 Appendix E Table 140 对应 v1.5 的 Appendix E Table 128。",
         size=11, color=GRAY, italic=True, first=True)


def s_hex_encoding(prs, layout):
    slide = content_slide(prs, layout, "值的编码：4 字符大写 HEX 字符串")
    tf = add_text(slide, BODY_X, 1.0, 6.1, 0.6)
    para(tf, "规范 §6.1.1：RDN 值 = UTF8String / PrintableString，内容为大写十六进制、恰好 4 字符、不省略前导零、无前缀。"
             "不是 2 字节二进制整数！", size=12, color=DARK, first=True)
    rows = [
        ["VID 0xFFF1", "\u201cFFF1\u201d", "✓ 正确"],
        ["PID 0x00B1", "\u201c00B1\u201d", "✓ 前导零必须保留"],
        ["PID 0x2A", "\u201c2A\u201d", "✗ 必须写 \u201c002A\u201d"],
        ["任意", "2 字节二进制 INTEGER", "✗ 解析必失败"],
    ]
    add_table(slide, BODY_X, 1.78, 6.2, ["值", "RDN 字符串", "判定"], rows,
              col_ratios=[1.6, 2.2, 3.0], body_size=10.5, row_h=0.44, header_h=0.34,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    add_code(slide, BODY_X, 4.42, 6.2, 1.95, [
        "// src/crypto/CHIPCryptoPAL.cpp",
        "inline constexpr size_t kVIDandPIDHexLength",
        "    = sizeof(uint16_t) * 2;                 // 必须 4 字节 ASCII",
        "VerifyOrReturnError(",
        "    attr.size() == kVIDandPIDHexLength,",
        "    CHIP_ERROR_WRONG_CERT_DN);",
        "Encoding::UppercaseHexToUint16(...);        // \u201cFFF1\u201d → 0xFFF1",
    ], size=10, title="SDK 解析侧代码")
    rx = BODY_X + 6.55
    add_box(slide, rx, 1.0, 5.75, 5.37, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, rx + 0.22, 1.18, 5.3, 5.1)
    para(tf, "为什么这个坑最隐蔽？", size=14, bold=True, color=RED, first=True, after=6)
    for t in [
        "OpenSSL / Wireshark 等通用工具按标准 X.509 显示，DN 一切\u201c看起来正常\u201d；",
        "Matter 的 ExtractVIDPIDFromX509Cert 却要求值恰为 4 字节 ASCII 大写 HEX；",
        "长度或大小写不符 → CHIP_ERROR_WRONG_CERT_DN → 证书被判格式非法（kDacFormatInvalid 303 / kPaiFormatInvalid 203）；",
        "排查时往往误以为是证书链问题，实际是 VID/PID 编码问题；",
        "工具链侧：chip-cert、openssl.cnf 模板、产测写入工具都要显式按字符串写入 VID/PID；",
        "自检办法：用 openssl x509 -noout -subject 打印，再对照 SDK 的解析规则逐字符核对。",
    ]:
        para(tf, t, size=12, bullet="▪", after=6)


def s_preferred_fallback(prs, layout):
    slide = content_slide(prs, layout, "两种编码方法：preferred vs fallback（§6.2.2.2）")
    add_box(slide, BODY_X, 1.02, 6.0, 2.35, fill=GREEN_BG)
    tf = add_text(slide, BODY_X + 0.22, 1.14, 5.6, 2.1)
    para(tf, "Preferred：用 Matter OID 作 RDN 属性类型", size=13.5, bold=True, color=GREEN, first=True, after=5)
    para(tf, "Subject: CN = Matter Test PAI,", size=11, font="Consolas", after=1)
    para(tf, "    1.3.6.1.4.1.37244.2.1 = FFF1,", size=11, font="Consolas", after=1)
    para(tf, "    1.3.6.1.4.1.37244.2.2 = 8000", size=11, font="Consolas", after=5)
    para(tf, "推荐方式：省 CN 空间、语义无歧义、便于 CA 做签发策略审计。", size=11.5, color=DARK)
    add_box(slide, BODY_X + 6.31, 1.02, 6.0, 2.35, fill=BLUE_BG)
    tf = add_text(slide, BODY_X + 6.53, 1.14, 5.6, 2.1)
    para(tf, "Fallback：塞进 commonName 字符串", size=13.5, bold=True, color=BLUE, first=True, after=5)
    para(tf, "Subject: CN = Matter Test DAC 0001", size=11, font="Consolas", after=1)
    para(tf, "    Mvid:FFF1 Mpid:8000", size=11, font="Consolas", after=5)
    para(tf, "给不支持的 CA 基础设施用：前缀后恰好 4 位大写 HEX；取最左合法匹配；有前缀但无完整合法值 → 整证书非法。",
         size=11.5, color=DARK)
    section_label(slide, "三条配套规则（规范与 SDK 行为一致）", x=BODY_X, y=3.52, w=6.5)
    rules = [
        "同一个字段（subject 或 issuer）内不得混用两种方法；",
        "字段里一旦出现 Matter OID 属性，就禁止再从 CN 解析 fallback —— SDK：OID 一无所获时才采用 CN 结果；",
        "DAC 的 issuer 必须与 PAI 的 subject 逐字节相同 → PAI 用了哪种编码，它签的 DAC 的 issuer 就固定是那种。",
    ]
    tf = add_text(slide, BODY_X, 3.95, BODY_W, 1.7)
    for i, t in enumerate(rules):
        para(tf, "%d. %s" % (i + 1, t), size=13, first=(i == 0), after=7)
    tf = add_text(slide, BODY_X, 5.85, BODY_W, 0.8)
    para(tf, "SDK 实现：ExtractVIDPIDFromX509Cert（CHIPCryptoPALmbedTLSCert.cpp / OpenSSL 版）遍历 Subject DN 的每个 RDN，"
             "依次匹配 commonName / matter-oid-vid / matter-oid-pid。", size=11.5, color=GRAY, italic=True, first=True)


def s_att_data(prs, layout):
    slide = content_slide(prs, layout, "Attestation 数据流：响应里到底有什么")
    add_box(slide, BODY_X, 1.02, 5.9, 0.72, fill=DARK)
    tf = add_text(slide, BODY_X, 1.1, 5.9, 0.55)
    para(tf, "AttestationResponse（规范 §11.18.6.2）", size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    add_box(slide, BODY_X, 1.92, 5.9, 1.55, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, BODY_X + 0.2, 2.0, 5.5, 1.4)
    para(tf, "AttestationElements [0]  ≤900B（RESP_MAX）", size=12.5, font="Consolas", bold=True, color=BLUE, first=True, after=3)
    para(tf, "AttestationSignature [1]  64B（ECDSA r||s）", size=12.5, font="Consolas", bold=True, color=RED, after=3)
    para(tf, "只有这两个字段 —— 没有证书！", size=12, bold=True, color=RED)
    add_box(slide, BODY_X, 3.75, 5.9, 2.55, fill=BLUE_BG)
    tf = add_text(slide, BODY_X + 0.2, 3.85, 5.5, 2.35)
    para(tf, "attestation-elements TLV（§11.18.4.6）", size=13, bold=True, color=BLUE, first=True, after=4)
    for t in [
        "certification_declaration [1] : OCTET STRING（CMS 签名的 CD）",
        "attestation_nonce [2] : 32B —— 原样带回请求里的 nonce",
        "timestamp [3] : uint32（epoch-s），给 DCL 查询用",
        "firmware_information [4] : 可选",
        "厂商字段必须用 fully-qualified tag，未知则忽略",
    ]:
        para(tf, t, size=11, font="Consolas", after=3)
    rx = BODY_X + 6.25
    section_label(slide, "TLV 解析的三条硬规则（SDK 强制）", x=rx, y=1.0, w=6.05)
    tf = add_text(slide, rx, 1.45, 6.05, 2.1)
    for i, t in enumerate([
        "第一个 context tag 必须是 1（CD）；",
        "后续 tag 必须严格递增（防混淆 / 防重放构造）；",
        "缺 CD / nonce / timestamp 任何一个 → CHIP_ERROR_MISSING_TLV_ELEMENT → kAttestationElementsMalformed (501)。",
    ]):
        para(tf, "%d. %s" % (i + 1, t), size=12.5, first=(i == 0), after=6)
    add_box(slide, rx, 3.6, 6.05, 1.4, fill=RED_BG)
    tf = add_text(slide, rx + 0.2, 3.72, 5.7, 1.2)
    para(tf, "attestation_challenge：16 字节", size=13, bold=True, color=RED, first=True, after=4)
    para(tf, "从当前 PASE / CASE 会话派生，不出现在任何链路载荷中（§11.18.4.7：challenge SHALL NOT be included）。",
         size=11.5, color=DARK)
    tf = add_text(slide, rx, 5.25, 6.05, 1.0)
    para(tf, "SDK 侧：src/credentials/DeviceAttestationConstructor.cpp 的 DeconstructAttestationElements()。",
         size=11, color=GRAY, italic=True, first=True)


def s_sign_design(prs, layout):
    slide = content_slide(prs, layout, "设备端签名：nonce 与 challenge 的双绑定设计")
    bx = BODY_X + 0.4
    b1 = add_box(slide, bx, 1.15, 3.4, 0.75, fill=BLUE_BG, line=BLUE)
    para(b1.text_frame, "attestation_elements (TLV)", size=12, font="Consolas", bold=True,
         color=BLUE, align=PP_ALIGN.CENTER, first=True)
    b2 = add_box(slide, bx + 4.0, 1.15, 3.4, 0.75, fill=RED_BG, line=RED)
    para(b2.text_frame, "attestation_challenge (16B, 会话保密)", size=11.5, font="Consolas", bold=True,
         color=RED, align=PP_ALIGN.CENTER, first=True)
    add_arrow(slide, bx + 1.7, 1.9, bx + 3.3, 2.5, color=DARK, w=1.75)
    add_arrow(slide, bx + 5.7, 1.9, bx + 4.5, 2.5, color=DARK, w=1.75)
    b3 = add_box(slide, bx + 2.2, 2.5, 3.4, 0.7, fill="F0F0F0", line="CCCCCC")
    para(b3.text_frame, "SHA-256（签名函数内部计算）", size=12, align=PP_ALIGN.CENTER, first=True)
    add_arrow(slide, bx + 3.9, 3.2, bx + 3.9, 3.62, color=DARK, w=1.75)
    b4 = add_box(slide, bx + 2.2, 3.62, 3.4, 0.7, fill=RED)
    para(b4.text_frame, "ECDSA-Sign(DAC 私钥, P-256)", size=12.5, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    add_arrow(slide, bx + 3.9, 4.32, bx + 3.9, 4.74, color=DARK, w=1.75)
    b5 = add_box(slide, bx + 2.2, 4.74, 3.4, 0.7, fill="F0F0F0", line="CCCCCC")
    para(b5.text_frame, "attestation_signature（64B r||s）", size=12, font="Consolas", align=PP_ALIGN.CENTER, first=True)
    rx = BODY_X + 8.15
    cards = [
        ("nonce：防重放", "每次配网唯一（CSPRNG 32B）；响应必须原样放回 elements[2] 并一起签名。重放旧响应 → nonce 对不上 → 502"),
        ("challenge：防挪用", "16B，绑定当前安全会话；中间人拿到签名也无法在另一条会话使用。缺 challenge → 无法防\u201c证据挪用\u201d"),
        ("timestamp：给 DCL 用", "Commissioner 按 DAC 的 notBefore 时间点查当时的 PAA 清单与吊销状态"),
    ]
    for i, (t, d) in enumerate(cards):
        y = 1.05 + i * 1.55
        add_box(slide, rx, y, 4.6, 1.4, fill="FAFAFA", line="DDDDDD")
        tf = add_text(slide, rx + 0.18, y + 0.08, 4.25, 0.32)
        para(tf, t, size=12.5, bold=True, color=RED, first=True)
        tf = add_text(slide, rx + 0.18, y + 0.4, 4.25, 0.95)
        para(tf, d, size=10.5, color=DARK, first=True)
    add_box(slide, BODY_X, 5.85, BODY_W, 1.0, fill=BLUE_BG)
    tf = add_text(slide, BODY_X + 0.25, 5.98, BODY_W - 0.5, 0.8)
    para(tf, "同一把 DAC 私钥的第二次出场：CSR 阶段 nocsr_tbs = nocsr_elements || attestation_challenge，同样的签名模型 ——"
             "私钥不出安全区、数据过来签名（§11.18.4.8）。", size=12.5, bold=True, color=BLUE, first=True)


def s_12steps_a(prs, layout):
    slide = content_slide(prs, layout, "Commissioner 的 12 步验证（上：1-6）")
    headers = ["步", "做什么", "SDK 关键动作", "失败错误码"]
    rows = [
        ["1", "入参完整性 + 大小检查", "各 buffer 非空；elements ≤ 900B（RESP_MAX）", "kInvalidArgument 701"],
        ["2", "PAI 必须存在", "paiDerBuffer 非空", "kPaiMissing 207"],
        ["3", "PAI / DAC 格式校验", "v3、ecdsa-with-SHA256、P-256、BC/KU 必须 critical、pathLen 规则、KU 位、SKID/AKID 在位", "kPaiFormatInvalid 203 / kDacFormatInvalid 303"],
        ["4", "VID/PID 交叉", "PAI 必须有 VID 且 == DAC.VID；DAC 必须有 PID；PAI 有 PID 则 == DAC.PID", "kDacVendorIdMismatch 305 / kDacProductIdMismatch 306"],
        ["5", "attestation 签名校验（用 DAC 公钥）", "SHA256(elements || challenge) 后 ECDSA 验签", "kAttestationSignatureInvalid 500 / kAttestationSignatureInvalidFormat 503"],
        ["6", "找 PAA", "取 PAI.AKID → 信任库按 SKID 查", "kPaaNotFound 101"],
    ]
    add_table(slide, BODY_X, 1.08, BODY_W, headers, rows,
              col_ratios=[0.5, 3.0, 6.3, 3.3], body_size=10.5, row_h=0.82, header_h=0.36,
              aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    tf = add_text(slide, BODY_X, 6.6, BODY_W, 0.5)
    para(tf, "顺序依据：DefaultDACVerifier::VerifyAttestationInformation()。规范允许优化执行顺序（§6.2.3.1），SDK 以\u201c最便宜且能尽早失败\u201d为原则。",
         size=11, color=GRAY, italic=True, first=True)


def s_12steps_b(prs, layout):
    slide = content_slide(prs, layout, "Commissioner 的 12 步验证（下：7-12）")
    headers = ["步", "做什么", "SDK 关键动作", "失败错误码"]
    rows = [
        ["7", "PAA 策略检查", "PAA 有 VID 则必须 == PAI.VID；PAA 禁止有 PID", "kPaiVendorIdMismatch 205 / kPaaFormatInvalid 105"],
        ["8", "DAC 有效期", "以签发时刻（notBefore）为基准检查（编译期可关 CURRENT_TIME_NOT_IMPLEMENTED）", "kDacExpired 300"],
        ["9", "证书链密码学验证", "ValidateCertificateChain(paa, pai, dac)，mbedTLS/OpenSSL 后端", "kDacSignatureInvalid 301 等（MapError）"],
        ["10", "elements TLV 解析 + nonce 比对", "第一个 tag=CD、tag 严格递增；nonce data_equal 比对", "kAttestationElementsMalformed 501 / kAttestationNonceMismatch 502"],
        ["11", "CD 校验：先验签、后语义交叉", "6 把钥匙验签 → Basic Info 基准 → 普通 / ODM 分支", "kCertificationDeclaration* 600~606"],
        ["12", "吊销检查（独立阶段）", "CheckForRevokedDACChain；无 delegate 打 WARNING 跳过", "kDacRevoked 302 / kPaiRevoked 202 / kPaiAndDacRevoked 208"],
    ]
    add_table(slide, BODY_X, 1.08, BODY_W, headers, rows,
              col_ratios=[0.5, 2.9, 5.9, 3.9], body_size=10.5, row_h=0.78, header_h=0.36,
              aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    add_box(slide, BODY_X, 6.22, BODY_W, 0.62, fill=RED_BG)
    tf = add_text(slide, BODY_X + 0.22, 6.32, BODY_W - 0.44, 0.45)
    para(tf, "注意第 5 步在前、第 9 步在后：先确认\u201c设备持有 DAC 私钥\u201d，再去验证证书链本身 —— 尽早失败，省掉昂贵的链验证。",
         size=11.5, bold=True, color=RED, first=True)


def s_code_verify(prs, layout):
    slide = content_slide(prs, layout, "第 5 步怎么做：ValidateAttestationSignature")
    add_code(slide, BODY_X, 1.05, 7.3, 4.6, [
        "CHIP_ERROR DeviceAttestationVerifier::ValidateAttestationSignature(",
        "        const P256PublicKey & pubkey,          // ← 从 DAC 证书提取的公钥",
        "        const ByteSpan & attestationElements,  // ← TLV 编码的 elements",
        "        const ByteSpan & attestationChallenge, // ← 会话 challenge (16B)",
        "        const P256ECDSASignature & signature)  // ← 设备返回的 64B",
        "{",
        "    Hash_SHA256_stream hashStream;",
        "    uint8_t md[kSHA256_Hash_Length];",
        "    MutableByteSpan messageDigestSpan(md);",
        "",
        "    hashStream.Begin();",
        "    hashStream.AddData(attestationElements);   // 先喂 elements",
        "    hashStream.AddData(attestationChallenge);  // 再拼 challenge",
        "    hashStream.Finish(messageDigestSpan);",
        "",
        "    return pubkey.ECDSA_validate_hash_signature(",
        "        messageDigestSpan.data(), messageDigestSpan.size(), signature);",
        "}",
    ], size=10.5, title="src/credentials/attestation_verifier/DeviceAttestationVerifier.cpp")
    rx = BODY_X + 7.65
    tf = add_text(slide, rx, 1.1, 4.65, 5.4)
    para(tf, "要点", size=14, bold=True, color=RED, first=True, after=8)
    for t in [
        "签名消息 = elements || challenge 的 SHA-256 哈希，与规范 §11.18.4.7 的 attestation_tbs 定义逐字对应；",
        "公钥取自 DAC 证书 —— \u201c证明持有私钥\u201d与\u201c证书链\u201d天然绑定；",
        "签名长度必须恰好 64B（r||s），否则报 503 InvalidFormat；",
        "验签失败 → 500，几乎总是设备端实现问题：拼接顺序错、没加 challenge、用错私钥；",
        "哈希是流式计算（Hash_SHA256_stream），elements 再大也只占固定内存；",
        "设备侧对应的接口是 SignWithDeviceAttestationKey —— 传入的 message 就是 attestation_tbs 本身。",
    ]:
        para(tf, t, size=12, bullet="▪", after=7)


def s_find_paa(prs, layout):
    slide = content_slide(prs, layout, "第 6 步：找 PAA —— AKID → SKID")
    add_code(slide, BODY_X, 1.05, 6.7, 2.5, [
        "class AttestationTrustStore {",
        "  public:",
        "    // 用 PAI 的 AKID（= PAA 的 SKID）查询信任库",
        "    virtual CHIP_ERROR GetProductAttestationAuthorityCert(",
        "        const ByteSpan & skid,",
        "        MutableByteSpan & outPaaDerBuffer) const = 0;",
        "};",
    ], size=10.5, title="src/credentials/attestation_verifier/DeviceAttestationVerifier.h")
    rows = [
        ["FileAttestationTrustStore", "chip-tool --paa-trust-store-path <dir>：从目录加载 PAA 文件，按 SKID 匹配"],
        ["TestAttestationTrustStore", "SDK 内置测试 PAA（如 Chip-Test-PAA-FFF1，SKID 6A:FD:22:77:...）；生产禁用"],
        ["自定义实现", "产品化 Commissioner 从 DCL 同步 PAA 清单 / 厂商私有信任库"],
    ]
    add_table(slide, BODY_X, 3.85, 6.7, ["信任库实现", "说明"], rows,
              col_ratios=[2.3, 4.4], body_size=10.5, row_h=0.62, header_h=0.34)
    rx = BODY_X + 7.05
    add_box(slide, rx, 1.05, 5.25, 1.7, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, rx + 0.2, 1.17, 4.9, 1.5)
    para(tf, "找不到时的日志与错误", size=13, bold=True, color=RED, first=True, after=4)
    para(tf, "Unable to find PAA, err: ..., PAI's AKID: 6A:FD:...", size=10, font="Consolas", after=3)
    para(tf, "→ kPaaNotFound (101)。排查方向：--paa-trust-store-path 目录不对 / PAA 未注册 DCL / AKID 与库内 SKID 不匹配。",
         size=11, color=DARK)
    add_box(slide, rx, 2.95, 5.25, 3.35, fill=BLUE_BG)
    tf = add_text(slide, rx + 0.2, 3.08, 4.9, 3.1)
    para(tf, "SKID / AKID 链式关系（真机日志可直接核对）", size=13, bold=True, color=BLUE, first=True, after=6)
    para(tf, "PAA  SKID=AKID  E9:16:0D:C4:...（自签）", size=10, font="Consolas", after=5)
    para(tf, "PAI  AKID = E9:16:0D:C4:...  = PAA.SKID", size=10, font="Consolas", after=5)
    para(tf, "DAC  AKID = EB:B4:9A:F1:...  = PAI.SKID", size=10, font="Consolas", after=6)
    para(tf, "首尾相接形成信任链；PAA 自签的标志就是 AKID == SKID。", size=11.5, color=DARK)


def s_cd_check(prs, layout):
    slide = content_slide(prs, layout, "第 11 步：CD 校验 —— 先签名、后语义")
    add_code(slide, BODY_X, 1.05, 7.2, 5.25, [
        "// ① 验 CMS 签名（ValidateCertificationDeclarationSignature）",
        "CMS_ExtractKeyId(cmsEnvelope, kid);        // 从 SignerInfo 取 KeyID",
        "mCdKeysTrustStore.LookupVerifyingKey(kid, verifyingKey);",
        "if (IsCdTestKey(kid) && !IsCdTestKeySupported())",
        "    return kCertificationDeclarationNoCertificateFound;",
        "CMS_Verify(cmsEnvelope, verifyingKey, cdPayload);",
        "",
        "// ② 语义交叉校验（ValidateCertificateDeclarationPayload）",
        "DecodeCertificationElements(...);          // TLV 解析",
        "VerifyOrReturnError(",
        "    cdContent.formatVersion == 1 &&",
        "    cdContent.certificationType < kReserved, ...);",
        "",
        "// 第一基准：Basic Information 集群",
        "VerifyOrReturnError(",
        "    cdContent.vendorId == deviceInfo.vendorId,   // ← Basic Info VID",
        "    kCertificationDeclarationInvalidVendorId);   // 604",
        "VerifyOrReturnError(",
        "    IsProductIdIn(cd, deviceInfo.productId),     // ← Basic Info PID",
        "    kCertificationDeclarationInvalidProductId);  // 605",
        "",
        "if (!cdContent.dacOriginVIDandPIDPresent) { /* 普通模式 */ }",
        "else                     { /* ODM 模式，见校验矩阵 */ }",
    ], size=10)
    rx = BODY_X + 7.55
    tf = add_text(slide, rx, 1.1, 4.75, 5.3)
    para(tf, "两个容易被忽略的强制点", size=14, bold=True, color=RED, first=True, after=8)
    para(tf, "1. PAI 带 PID 时，这个 PID 也必须在 CD 的 product_id_array 里（普通模式），或等于 dac_origin_product_id（ODM 模式）。"
             "\u201cPAI 通配所有 PID\u201d的前提是它根本不带 PID。", size=12, bullet="▪", after=8)
    para(tf, "2. 即使 ODM 模式，CD.vendor_id / product_id_array 仍然必须匹配 Basic Info —— "
             "设备 Basic Info 报品牌身份，DAC 报原厂身份，CD 用 dac_origin_* 把两者绑定。", size=12, bullet="▪", after=8)
    para(tf, "错误码一览：601 找不到验签钥匙 / 602 签名错 / 603 格式错 / 604 VID / 605 PID / 606 PAA 不在白名单。",
         size=12, bullet="▪", after=8)
    para(tf, "deviceInfo.vendorId / productId 的注释：\u201creported by device in Basic Information cluster\u201d。",
         size=11, color=GRAY, italic=True)


def s_cd_what(prs, layout):
    slide = content_slide(prs, layout, "CD 是什么 + CMS 结构（§6.3.1）")
    tf = add_text(slide, BODY_X, 0.98, BODY_W, 0.75)
    para(tf, "CD = CSA 联盟（不是厂商！）签发的一段 CMS 签名数据，声明 \u201cVID=V 的这些 PID 已通过认证（类型 T）\u201d。"
             "产品认证通过后颁发，随固件/工厂数据写入设备，配网时在 AttestationResponse 里上交。", size=12.5, first=True, after=3)
    para(tf, "规范明文：Certification Declarations SHALL NOT be generated by any Node —— 设备只能存储并转发。",
         size=12, bold=True, color=RED)
    add_box(slide, BODY_X, 1.95, 6.0, 0.55, fill=DARK)
    tf = add_text(slide, BODY_X, 2.0, 6.0, 0.45)
    para(tf, "CertificationDeclaration ::= SEQUENCE", size=12, font="Consolas", bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
    cms = [
        ("version  = v3(3)", "F0F0F0", None),
        ("digestAlgorithm = sha256", "F0F0F0", None),
        ("encapContentInfo", "E8E8E8", None),
        ("   eContentType = pkcs7-data", "FAFAFA", None),
        ("   eContent = OCTET STRING ← TLV 明文（下页）", "FAFAFA", None),
        ("signerInfo", "E8E8E8", None),
        ("   subjectKeyIdentifier ← KeyID，选验签公钥的依据", BLUE_BG, BLUE),
        ("   digestAlgorithm = sha256", "FAFAFA", None),
        ("   signatureAlgorithm = ecdsa-with-SHA256", "FAFAFA", None),
        ("   signature = OCTET STRING", "FAFAFA", None),
    ]
    y = 2.62
    for t, bg, fg in cms:
        add_box(slide, BODY_X + 0.3, y, 5.7, 0.36, fill=bg, line=None, radius=0.15)
        tf = add_text(slide, BODY_X + 0.45, y + 0.035, 5.5, 0.3)
        para(tf, t, size=10.5, font="Consolas", bold=(fg is not None),
             color=(fg or DARK), first=True)
        y += 0.4
    rx = BODY_X + 6.6
    add_box(slide, rx, 1.95, 5.7, 4.35, fill="FAFAFA", line="DDDDDD")
    tf = add_text(slide, rx + 0.22, 2.1, 5.25, 4.1)
    para(tf, "要点", size=14, bold=True, color=RED, first=True, after=6)
    for t in [
        "CMS 是 RFC 5652 SignedData 的最小子集 —— 没有 certificates / crls 字段；",
        "验签公钥不随 CD 传输：Commissioner 用 signerInfo 里的 SKID（20B）在内置钥匙表里查（下页）；",
        "digest 与签名算法固定 sha256 + ecdsa-with-SHA256，曲线 secp256r1；",
        "eContent 里是 TLV 明文 —— 验签通过后取出再做语义交叉校验；",
        "CD 与证书链相互独立：换 CD 不需要换 DAC，但三处 VID/PID 语义必须一致；",
        "CSA 可轮换/新增签名钥匙（002~005 即后补），Commissioner 按内置清单更新。",
    ]:
        para(tf, t, size=12, bullet="▪", after=7)


def s_cd_tlv(prs, layout):
    slide = content_slide(prs, layout, "CD 内嵌 TLV 明文：certification-elements（修正版）")
    headers = ["tag", "字段", "类型", "说明"]
    rows = [
        ["0", "format_version", "uint16", "固定 = 1"],
        ["1", "vendor_id", "uint16", "必须 == Basic Info 的 VendorID"],
        ["2", "product_id_array", "uint16[1..100]", "必须 ⊇ Basic Info 的 ProductID"],
        ["3", "device_type_id", "uint32", "设备主设备类型"],
        ["4", "certificate_id", "string[19]", "CSA 分配的全局唯一编号，如 \u201cZIG20142ZB330003-24\u201d"],
        ["5", "security_level", "uint8", "★ 8-bit（旧版误写 32-bit）；保留，置 0，读时忽略"],
        ["6", "security_information", "uint16", "保留，置 0，读时忽略"],
        ["7", "version_number", "uint16", "CD 自身版本号（联盟分配，与固件版本无关）"],
        ["8", "certification_type", "uint8", "0=开发测试  1=临时(Provisional)  2=正式；≥3 保留，SDK 直接拒绝"],
        ["9", "dac_origin_vendor_id", "uint16", "可选；与 [10] 必须成对出现 / 成对缺省"],
        ["10", "dac_origin_product_id", "uint16", "可选；ODM 场景指向原厂 DAC 的 PID"],
        ["11", "authorized_paa_list", "OCTET STRING[20] ×1..10", "可选；允许签发 PAI 的 PAA SKI 白名单"],
    ]
    add_table(slide, BODY_X, 1.05, BODY_W, headers, rows,
              col_ratios=[0.6, 3.0, 2.6, 6.0], body_size=10.5, row_h=0.415, header_h=0.34,
              aligns=[PP_ALIGN.CENTER, PP_ALIGN.LEFT, PP_ALIGN.LEFT, PP_ALIGN.LEFT])
    tf = add_text(slide, BODY_X, 6.65, BODY_W, 0.4)
    para(tf, "未列出的 context tag：保留，Commissioner 必须静默忽略（SDK 的解析器按 tag-order 严格递增读取）。",
         size=11, color=GRAY, italic=True, first=True)


def s_cd_keys(prs, layout):
    slide = content_slide(prs, layout, "CD 验签公钥：SDK 内置 6 把钥匙（gCdSigningKeys[]）")
    headers = ["用途", "Key ID（SKID, 20 字节）"]
    rows = [
        ["测试 CD Signing Key（仅开发）", "62:FA:82:33:59:AC:FA:A9:96:3E:1C:FA:14:0A:DD:F5:04:F3:71:60"],
        ["官方 Signing Key 001", "FE:34:3F:95:99:47:76:3B:61:EE:45:39:13:13:38:49:4F:E6:7D:8E"],
        ["官方 Signing Key 002", "DD:04:DB:58:5B:21:4C:1C:58:15:87:E6:56:8D:F4:87:B6:DD:C7:01"],
        ["官方 Signing Key 003", "47:10:35:E7:C0:4E:AA:A8:BE:7C:4D:4C:13:E3:E4:C2:09:95:A8:4B"],
        ["官方 Signing Key 004", "F6:86:03:A3:69:2E:98:10:72:41:9E:A1:E1:AB:38:54:BD:77:95:D3"],
        ["官方 Signing Key 005", "63:7F:26:34:AD:62:EA:FE:6A:F6:62:EF:B9:6F:6F:D2:FC:BF:FC:2F"],
    ]
    add_table(slide, BODY_X, 1.08, BODY_W, headers, rows,
              col_ratios=[3.2, 7.4], body_size=11, row_h=0.44, header_h=0.36)
    facts = [
        ("信任锚", "官方钥匙的证书必须链到 CSA 根：CN = Matter Certification and Testing CA, O = CSA（SDK 内嵌 gCdRootCert）；运行时 AddTrustedKey(derCert) 追加的钥匙会先做链验证"),
        ("测试钥匙开关", "EnableCdTestKeySupport(bool)：默认 true；生产必须设 false，否则测试 CD 仍然通行（日志 Disallowing CD signed by test key）"),
        ("日志对照", "chip-tool 打印 CD signing key identifier: FE:34:... 即命中官方 001 号钥匙"),
    ]
    for i, (t, d) in enumerate(facts):
        y = 4.25 + i * 0.88
        add_box(slide, BODY_X, y, BODY_W, 0.76, fill="FAFAFA" if i % 2 == 0 else WHITE, line="DDDDDD")
        tf = add_text(slide, BODY_X + 0.2, y, 1.7, 0.76, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, t, size=12, bold=True, color=RED, first=True)
        tf = add_text(slide, BODY_X + 2.0, y, BODY_W - 2.3, 0.76, anchor=MSO_ANCHOR.MIDDLE)
        para(tf, d, size=11, color=DARK, first=True)


def s_matrix(prs, layout):
    slide = content_slide(prs, layout, "VID/PID 交叉校验总矩阵（背下这张表）")
    headers = ["校验关系", "普通模式", "ODM 模式（CD 含 dac_origin_*）"]
    rows = [
        ["CD.vendor_id ↔ Basic Info VendorID", "必须", "必须（不变）"],
        ["CD.product_id_array ⊇ Basic Info ProductID", "必须", "必须（不变）"],
        ["DAC.VID ↔ CD.vendor_id", "必须", "✗（换成下行）"],
        ["DAC.VID ↔ CD.dac_origin_vendor_id", "—", "必须"],
        ["PAI.VID ↔ CD.vendor_id", "必须", "✗（换成 dac_origin_vendor_id）"],
        ["DAC.PID ↔ CD", "∈ product_id_array", "== dac_origin_product_id"],
        ["PAI.PID（若有）↔ CD", "∈ product_id_array", "== dac_origin_product_id"],
        ["DAC.VID ↔ PAI.VID（第 4 步，与 CD 无关）", "必须", "必须"],
        ["PAA.SKID ∈ authorized_paa_list（若存在）", "必须", "必须"],
    ]
    add_table(slide, BODY_X, 1.08, BODY_W, headers, rows,
              col_ratios=[5.2, 3.0, 4.2], body_size=11, row_h=0.475, header_h=0.38,
              aligns=[PP_ALIGN.LEFT, PP_ALIGN.CENTER, PP_ALIGN.CENTER])
    tf = add_text(slide, BODY_X, 6.35, BODY_W, 0.7)
    para(tf, "读法：ODM 模式只替换\u201c证书身份 ↔ CD\u201d这一组关系；\u201cCD ↔ Basic Info\u201d与\u201c证书链内部\u201d的关系任何模式下都成立。",
         size=12, bold=True, color=RED, first=True)
    para(tf, "真机实例：DAC 0x1470/0x8006 ↔ dac_origin_*；Basic Info 0x149A/0x3005 ↔ CD 常规字段（见下一页真机日志）。",
         size=11.5, color=GRAY)


def s_faq_a(prs, layout):
    slide = content_slide(prs, layout, "FAQ（上）：PID 与 CD 的三种典型疑问")
    faqs = [
        ("Q1  PAI 能不能带\u201c一组 PID\u201d，让 DAC 报列表里任意一个都通过？",
         "不能。PAI 的 PID 是 X.509 DN 里的单值 RDN，只能 0 个或 1 个。要覆盖多型号只有一种姿势：PAI 不带 PID（对该 VID 下所有 PID 都有签发权）+ 所有已认证 PID 进 CD 的 product_id_array。"
         "SDK 第 4 步只在 PAI 带 PID 时才强制 PAI.PID == DAC.PID；第 11 步强制 DAC.PID ∈ CD 列表，且 PAI.PID（若有）也 ∈ CD 列表。", 1.75),
        ("Q2  DAC 的 PID 不在 CD 的 product_id_array 里，能入网吗？",
         "不能。规范 §6.2.3.1 是 SHALL 级校验，SDK 返回 kCertificationDeclarationInvalidProductId (605)，chip-tool 直接终止 commissioning。VID 不匹配对应 604。唯一例外见下一页 Q3/Q4。", 1.5),
        ("Q3  真有一台\u201cPID 对不上却入网成功\u201d的设备，怎么解释？",
         "① ODM 模式：DAC 匹配的是 CD 的 dac_origin_* 字段（日志会打印 DAC origin VID/PID）—— 合规；② Commissioner 实现有缺陷 / 开发模式放宽；③ 某些 Commissioner 不读 Basic Info，少了一道交叉；"
         "④ 吊销 / 固件信息检查允许降级，但不影响 VID/PID 强校验。结论：对量产零售产品，不合理，除非命中 ①。", 1.95),
    ]
    y = 1.02
    for q, a, h in faqs:
        add_box(slide, BODY_X, y, BODY_W, h, fill="FAFAFA", line="DDDDDD")
        tf = add_text(slide, BODY_X + 0.22, y + 0.07, BODY_W - 0.44, 0.38)
        para(tf, q, size=13.5, bold=True, color=RED, first=True)
        tf = add_text(slide, BODY_X + 0.22, y + 0.44, BODY_W - 0.44, h - 0.52)
        para(tf, a, size=11.5, color=DARK, first=True, line=1.12)
        y += h + 0.1


def s_faq_b(prs, layout):
    slide = content_slide(prs, layout, "FAQ（下）：dac_origin 与量产实践")
    faqs = [
        ("Q4  dac_origin_vendor_id / dac_origin_product_id 是什么（\u201cOriginal PID\u201d）？",
         "规范里没有叫 Original PID 的字段，正式名称就是 CD 的 dac_origin_* [9]/[10]：必须成对出现或成对缺省；为 ODM 贴牌设计 —— ODM 出厂时 DAC/PAI 已固化自己的 VID/PID，"
         "品牌方写入自己的 CD（常规字段=品牌 + dac_origin_*=原厂链）。origin 值与常规值可以相同。关键洞察：Basic Info 报品牌身份，DAC 报原厂身份，CD 把两套身份绑定。", 1.55),
        ("Q5  Basic Information 集群报的 VID/PID 和证书不一致会怎样？",
         "CD 校验直接失败（604/605）：CD.vendor_id 必须 == Basic Info VendorID、product_id_array 必须 ⊇ Basic Info ProductID。烧录工厂数据时，证书、CD、Basic Info 三处必须同源配置。", 1.05),
        ("Q6  签名为什么要同时掺 nonce 和 challenge？",
         "nonce 防重放（每次配网唯一，响应必须原样带回并签名）；challenge 绑定安全会话（16B，PASE/CASE 派生，从不上链路），防中间人把签名挪到别的会话使用 —— 既防旧证据重放，也防证据挪用。", 1.15),
        ("Q7  开发用 0xFFF1/0x8000 测试证书，量产要注意什么？",
         "换 CSA 官方/自有 PAA（注册 DCL）；换联盟签发的正式 CD（type=1/2）；EnableCdTestKeySupport(false)；真实 PKI 签发 DAC，私钥进安全元件；实现 revocation delegate。", 1.15),
    ]
    y = 1.02
    for q, a, h in faqs:
        add_box(slide, BODY_X, y, BODY_W, h, fill="FAFAFA", line="DDDDDD")
        tf = add_text(slide, BODY_X + 0.22, y + 0.05, BODY_W - 0.44, 0.36)
        para(tf, q, size=13, bold=True, color=RED, first=True)
        tf = add_text(slide, BODY_X + 0.22, y + 0.4, BODY_W - 0.44, h - 0.48)
        para(tf, a, size=11, color=DARK, first=True, line=1.1)
        y += h + 0.1


def s_device_impl(prs, layout):
    slide = content_slide(prs, layout, "设备端落地（Silicon Labs）")
    add_code(slide, BODY_X, 1.02, 6.3, 2.6, [
        "class DeviceAttestationCredentialsProvider {",
        "  virtual CHIP_ERROR GetCertificationDeclaration(...);   // CD",
        "  virtual CHIP_ERROR GetFirmwareInformation(...);        // 可为空",
        "  virtual CHIP_ERROR GetDeviceAttestationCert(...);      // DAC DER",
        "  virtual CHIP_ERROR GetProductAttestationIntermediateCert(...);",
        "  // 签名在设备内部完成，私钥永远不返回：",
        "  virtual CHIP_ERROR SignWithDeviceAttestationKey(",
        "      const ByteSpan & messageToSign,",
        "      MutableByteSpan & outSignatureBuffer) = 0;",
        "};",
    ], size=9.5, title="src/credentials/DeviceAttestationCredsProvider.h")
    tf = add_text(slide, BODY_X, 3.75, 6.3, 2.6)
    para(tf, "两个现成实现", size=13, bold=True, color=RED, first=True, after=4)
    para(tf, "开发：ExampleDACProvider（src/credentials/examples/）内置 DevelopmentCerts 的 DAC/PAI/CD，软算签名；",
         size=11.5, bullet="▪", after=4)
    para(tf, "量产：DAC/PAI/CD 从工厂数据区读取，SignWithDeviceAttestationKey 下发安全元件（SE/TEE/TA）内部完成；",
         size=11.5, bullet="▪", after=4)
    para(tf, "设计要点：NVM3 只存\u201c凭据区基址 + 各段偏移/长度\u201d，证书本体放独立 flash 凭据区，DAC 私钥放 SE/TA 不可导出。",
         size=11.5, bullet="▪", after=4)
    para(tf, "读取工具：commander nvm3 read / parse。", size=11.5, bullet="▪")
    rx = BODY_X + 6.65
    headers = ["NVM3 Key", "值", "内容"]
    rows = [
        ["Creds_KeyId", "0x087220", "CD 签名钥匙序号"],
        ["Creds_Base_Addr", "0x087221", "凭据区 flash 基址"],
        ["Creds_DAC_Offset / Size", "0x087222/23", "DAC 证书位置"],
        ["Creds_PAI_Offset / Size", "0x087224/25", "PAI 证书位置"],
        ["Creds_CD_Offset / Size", "0x087226/27", "CD 位置"],
    ]
    add_table(slide, rx, 1.02, 5.65, headers, rows,
              col_ratios=[2.4, 1.5, 2.2], body_size=10, row_h=0.4, header_h=0.34)
    tf = add_text(slide, rx, 3.5, 5.65, 1.3)
    para(tf, "Key = 0x087000 (Matter NVM3 域) | (0x2 << 8) | id，合法区间 [0x087200, 0x087FFF]。",
         size=11, first=True, after=3)
    para(tf, "量产工厂数据清单：DAC / PAI / CD / DAC 私钥（SE 内）/ Basic Info 数据（VID、PID、序列号等，与证书同源）/ SPAKE2+ 参数。",
         size=11, bold=True, color=BLUE)
    add_box(slide, rx, 5.0, 5.65, 1.35, fill=GREEN_BG)
    tf = add_text(slide, rx + 0.2, 5.1, 5.25, 1.15)
    para(tf, "产线提示", size=12, bold=True, color=GREEN, first=True, after=3)
    para(tf, "首次量产先小批量跑完整 chip-tool 配网 + 验签回归；写凭据时按 Key 表逐项核对长度（DER ≤ 600B）。", size=11, color=DARK)


def s_errors(prs, layout):
    slide = content_slide(prs, layout, "错误码速查与排查")
    headers = ["错误码", "含义", "排查方向"]
    rows = [
        ["101 kPaaNotFound", "信任库找不到 PAA", "--paa-trust-store-path 目录不对 / 未注册 DCL / AKID↔SKID 不匹配"],
        ["203/303 FormatInvalid", "PAI / DAC 格式非法", "BC、KU 非 critical；pathLen 错；KU 位错；OID 用错 arc"],
        ["305/306 Vid/PidMismatch", "DAC 与 PAI 的 VID/PID 不一致", "产线证书配置；PAI 缺 VID；DAC 缺 PID"],
        ["301 kDacSignatureInvalid", "证书链签名验证失败", "PAI 不是签发该 DAC 的钥匙（常为产线刷错 PAI）"],
        ["500 AttestationSignatureInvalid", "attestation 签名失败", "设备端：拼接顺序错、没加 challenge、用错私钥"],
        ["502 NonceMismatch", "nonce 不匹配", "设备没把请求里的 nonce 原样放进 elements[2]"],
        ["604/605 InvalidVendor/ProductId", "CD 与 VID/PID 交叉失败", "证书 ↔ CD ↔ Basic Info 三处不同源（FAQ-Q5）"],
        ["601 + Disallowing test key", "CD 用测试钥匙且被关闭", "量产设备 + Commissioner 已 EnableCdTestKeySupport(false)"],
    ]
    add_table(slide, BODY_X, 1.08, BODY_W, headers, rows,
              col_ratios=[3.3, 3.0, 6.2], body_size=10.5, row_h=0.575, header_h=0.36)
    tf = add_text(slide, BODY_X, 6.35, BODY_W, 0.7)
    para(tf, "完整枚举见 SDK src/credentials/attestation_verifier/DeviceAttestationVerifier.h（100/200/300/400/500/600/700 分段）；"
             "文字描述用 GetAttestationResultDescription()。", size=11, color=GRAY, italic=True, first=True)


def s_checklist(prs, layout):
    slide = content_slide(prs, layout, "生产 Checklist")
    items = [
        "证书：v3、ecdsa-with-SHA256、P-256；BC/KU 均 critical；DAC KU 仅 digitalSignature；PAI/PAA KU 为 keyCertSign+cRLSign；PAI pathLen=0；PAA pathLen 缺省或 1；DER ≤ 600B",
        "VID/PID 编码：preferred（OID .2.1/.2.2）或 fallback（Mvid:/Mpid:）二选一；值是 4 字符大写 HEX 字符串",
        "链条：恰好 PAA→PAI→DAC 三层；AKID/SKID 首尾相接；DAC.issuer 与 PAI.subject 逐字节相同",
        "三处身份同源：DAC/PAI ↔ CD ↔ Basic Information 集群",
        "ODM 贴牌：CD 成对携带 dac_origin_*，品牌身份进常规字段",
        "Commissioner：EnableCdTestKeySupport(false)；PAA 信任库来自 DCL；实现 revocation delegate",
        "设备：DAC 私钥仅存安全元件；签名走 SignWithDeviceAttestationKey",
        "回归：用规范附录 F 测试向量做验签单测；chip-tool 全流程 + 异常注入（错 PID / 错 nonce / 换 CD）验证错误路径",
    ]
    for i, t in enumerate(items):
        y = 1.08 + i * 0.62
        add_box(slide, BODY_X, y + 0.01, 0.34, 0.34, fill=GREEN, radius=0.5)
        tfn = add_text(slide, BODY_X, y + 0.05, 0.34, 0.28)
        para(tfn, "✓", size=12, bold=True, color=WHITE, align=PP_ALIGN.CENTER, first=True)
        tf = add_text(slide, BODY_X + 0.52, y - 0.02, BODY_W - 0.52, 0.58)
        para(tf, t, size=12, first=True)
    tf = add_text(slide, BODY_X, 6.35, BODY_W, 0.7)
    para(tf, "参考：Matter Core Spec R1.5 §6.1-6.3、§11.18、Appendix E/F；SDK src/credentials/ 与 src/platform/silabs/；DCL：dcl.csa-iot.org；Matter Handbook（handbook.buildwithmatter.com）",
         size=11, color=GRAY, italic=True, first=True)


# ============================================================
# v2.1 新增：配网全景 / 文档层级 / 真机日志
# ============================================================

IMG_DIR = r"D:\hrf\h\files\matter\handbook-img"


def s_commissioning_overview(prs, layout):
    slide = content_slide(prs, layout, "配网全流程全景（官方手册 + 真机日志）")
    pic = os.path.join(IMG_DIR, "primer-commissioning.png")
    if os.path.exists(pic):
        slide.shapes.add_picture(pic, Inches(2.02), Inches(1.0), width=Inches(9.3))
    stages_l = [
        ("1", "BLE 扫描连接", "Discriminator 匹配广播"),
        ("2", "SecurePairing", "PASE / SPAKE2+ 会话"),
        ("3", "ReadCommissioningInfo", "Basic Info 自报 VID/PID"),
        ("4", "ArmFailSafe / ConfigRegulatory", "60s 失败保险 / 地区配置"),
        ("5", "请求 PAI 证书", "CertChainRequest type=2，470B"),
        ("6", "请求 DAC 证书", "CertChainRequest type=1，481B"),
        ("7", "设备认证", "nonce 32B → elements 423B + 签名 64B"),
        ("8", "AttestationVerification ★", "本档 12 步验证"),
    ]
    stages_r = [
        ("9", "AttestationRevocationCheck", "查 DCL 吊销（可跳过）"),
        ("10", "NOCSR → GenerateNOCChain", "复用 DAC 私钥签名"),
        ("11", "SendTrustedRootCert → SendNOC", "装根证书 + 运营证书"),
        ("12", "Thread 网络配置与启用", "AddOrUpdate → ConnectNetwork"),
        ("13", "DNS-SD 发现 → CASE 建立", "Sigma1/2/3，用 NOC 建会话"),
        ("14", "CommissioningComplete", "成功后自动撤失败保险"),
        ("15", "Cleanup", "配网完成，断开 BLE"),
    ]
    for col, items in ((0, stages_l), (1, stages_r)):
        x = BODY_X + col * 6.28
        for i, (num, t, d) in enumerate(items):
            y = 4.05 + i * 0.375
            tfn = add_text(slide, x, y, 0.35, 0.3)
            para(tfn, num, size=11, bold=True, color=RED, first=True)
            tft = add_text(slide, x + 0.38, y, 2.35, 0.3)
            para(tft, t, size=10.5, bold=True, color=DARK, first=True)
            tfd = add_text(slide, x + 2.78, y, 3.4, 0.3)
            para(tfd, d, size=10, color=GRAY, first=True)
    tf = add_text(slide, BODY_X, 6.95, BODY_W, 0.34)
    para(tf, "阶段名与真机日志一一对应（commissioning-raspi-log.md）；★ 为设备认证相关。图源：Matter Handbook《Commissioning》",
         size=10, color=GRAY, italic=True, first=True)


def s_doc_hierarchy(prs, layout):
    slide = content_slide(prs, layout, "认证文档层级总览（官方手册）")
    pic = os.path.join(IMG_DIR, "primer-attestation-document-hierarchy.png")
    if os.path.exists(pic):
        slide.shapes.add_picture(pic, Inches(7.5), Inches(0.98), height=Inches(5.4))
    items = [
        ("Attestation Information", "= AttestationElements + AttestationSignature —— 即 AttestationResponse 的两个已知字段"),
        ("Attestation TBS", "= attestation_elements || attestation_challenge —— 签名消息本体"),
        ("Attestation Elements TLV", "= CD[1] + nonce[2] + timestamp[3] + firmware_info[4]（可选）+ 厂商字段"),
        ("ECDSA Sign", "用 DAC 私钥在设备内部签名（私钥不出安全区）"),
        ("Attestation Challenge", "PASE/CASE 会话派生 16B，不出现在任何链路载荷"),
        ("Firmware Information（可选）", "CD 版本号 + 固件组件摘要，可对照 DCL（Commissioner 可不支持）"),
    ]
    y = 1.1
    for t, d in items:
        add_box(slide, BODY_X, y, 6.7, 0.78, fill="FAFAFA" if int(y * 100) % 2 else "FAFAFA", line="E2E2E2")
        tf = add_text(slide, BODY_X + 0.18, y + 0.05, 6.35, 0.3)
        para(tf, t, size=12, bold=True, color=BLUE, first=True)
        tf = add_text(slide, BODY_X + 0.18, y + 0.36, 6.35, 0.38)
        para(tf, d, size=10.5, color=DARK, first=True)
        y += 0.87
    tf = add_text(slide, BODY_X, 6.5, 6.7, 0.5)
    para(tf, "图源：Matter Handbook《Attestation》（2025-10-14）", size=10, color=GRAY, italic=True, first=True)


def s_real_log(prs, layout):
    slide = content_slide(prs, layout, "真机日志解读：量产设备 + ODM 模式实例")
    add_code(slide, BODY_X, 1.02, 6.55, 5.3, [
        "[SVR] OnReadCommissioningInfo - vendorId=0x149A productId=0x3005",
        "[-] --> DAC's VID: 0x1470, PID: 0x8006",
        "[-] --> DAC certificate AKID: EB:B4:9A:F1:...:D0:BF:9F:9A",
        "[-] --> PAI certificate SKID: EB:B4:9A:F1:...   = DAC.AKID  OK",
        "[-] --> PAI certificate AKID: E9:16:0D:C4:...:F3:12:22",
        "[-] --> PAA certificate SKID: E9:16:0D:C4:...   = PAI.AKID  OK",
        "[-] --> PAA certificate AKID: E9:16:0D:C4:...   =SKID 自签 OK",
        "[-] CD signing key identifier: FE:34:3F:95:...  官方 001",
        "[-] --> VID: 0x149A",
        "[-] --> Device type ID: 0x0000_0202",
        "[-] --> Certification type: 2 (Certified device)",
        "[-] --> DAC origin VID: 0x1470, PID: 0x8006",
        "[CTL] Successfully finished 'AttestationVerification'",
        "[-] WARNING: No revocation delegate available...",
    ], size=9.5, title="commissioning-raspi-log.md（HOPERF 量产设备，节选）")
    rx = BODY_X + 6.9
    headers = ["来源", "VID / PID", "校验对象"]
    rows = [
        ["DAC / PAI 证书 DN", "0x1470 / 0x8006", "CD.dac_origin_*"],
        ["CD 常规字段", "0x149A / 数组", "Basic Info 集群"],
        ["Basic Info 集群", "0x149A / 0x3005", "CD.vendor_id / array"],
        ["CD.dac_origin_*", "0x1470 / 0x8006", "DAC/PAI 证书 DN"],
    ]
    add_table(slide, rx, 1.02, 5.4, headers, rows,
              col_ratios=[2.0, 1.9, 2.0], body_size=10, row_h=0.44, header_h=0.34)
    add_box(slide, rx, 3.6, 5.4, 1.55, fill=RED_BG)
    tf = add_text(slide, rx + 0.2, 3.72, 5.0, 1.35)
    para(tf, "读法", size=13, bold=True, color=RED, first=True, after=3)
    para(tf, "设备自报品牌身份（0x149A/0x3005），证书携带原厂身份（0x1470/0x8006），CD 用 dac_origin_* 把两者绑定 —— 校验矩阵右列的完整落地。",
         size=11.5, color=DARK)
    add_box(slide, rx, 5.3, 5.4, 1.02, fill=BLUE_BG)
    tf = add_text(slide, rx + 0.2, 5.4, 5.0, 0.85)
    para(tf, "certification_type: 2 (Certified device) —— 正式认证；PAA/PAI 有效期 100 年；CD 命中官方 001 号钥匙。",
         size=11.5, color=DARK, first=True)


# ============================================================
# 主流程
# ============================================================

def main():
    prs = Presentation(SRC)
    cover = prs.slides[0]
    template_slide = prs.slides[1]
    thanks = prs.slides[2]
    layout = template_slide.slide_layout

    # --- 封面副标题更新（只改文本，保留样式基调）---
    for shp in cover.shapes:
        if shp.has_text_frame and "PAA PAI DAC" in shp.text_frame.text:
            shp.width = Inches(5.85)
            shp.left = Inches(7.40)
            tf = shp.text_frame
            p = tf.paragraphs[0]
            for r in list(p.runs)[1:]:
                r._r.getparent().remove(r._r)
            if p.runs:
                p.runs[0].text = "PAA PAI DAC + CD | 验签详解"
                style_run(p.runs[0], size=28, bold=True, color=DARK)
            else:
                run = p.add_run()
                run.text = "PAA PAI DAC + CD | 验签详解"
                style_run(run, size=28, bold=True, color=DARK)

    # --- 依次生成内容页（追加到末尾，稍后统一排序）---
    builders = [
        s_toc, s_revisions, s_new_contents, s_commissioning_overview,
        s_flow_position, s_separate,
        s_analogy, s_pki, s_vidpid_map, s_cert_table, s_oid,
        s_hex_encoding, s_preferred_fallback, s_att_data, s_doc_hierarchy,
        s_sign_design,
        s_12steps_a, s_12steps_b, s_code_verify, s_find_paa, s_cd_check,
        s_cd_what, s_cd_tlv, s_cd_keys, s_matrix, s_real_log, s_faq_a, s_faq_b,
        s_device_impl, s_errors, s_checklist,
    ]
    for b in builders:
        b(prs, layout)

    # --- 重排：删除空白模板页，THANKS 移到最后 ---
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    blank_el, thanks_el = ids[1], ids[2]
    rId = blank_el.get(qn("r:id"))
    sldIdLst.remove(blank_el)
    try:
        prs.part.drop_rel(rId)
    except Exception:
        pass
    sldIdLst.remove(thanks_el)
    sldIdLst.append(thanks_el)

    prs.save(DST)

    # --- 快速自检 ---
    check = Presentation(DST)
    print("slides:", len(check.slides._sldIdLst))
    for i, s in enumerate(check.slides, 1):
        t = s.shapes.title.text if s.shapes.title is not None else "(no title)"
        print(f"  {i:02d}. {t}")


if __name__ == "__main__":
    main()
